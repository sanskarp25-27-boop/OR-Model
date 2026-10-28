"""The Primary Sales record's downloadable Excel report (the "Download report" button in the OR app).

Five sheets built from ui_api.tracker_report: Summary, Secondary Sales, Stock Ledger, Accounting Ledger and
Rate Difference Notes. Formatted for reading: title banners, figure tiles, styled tables with banded rows,
coloured statuses, gains in green and losses in red, links to the ERPNext documents, filters and frozen headers.
"""

from collections import defaultdict
from io import BytesIO

import frappe
from frappe.utils import flt, get_url_to_form, getdate
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# palette (matches the OR app)
INDIGO = "4338CA"
INDIGO_DARK = "312E81"
LAVENDER = "EEF0FF"
SECTION = "E0E7FF"
BAND = "F8F9FD"
LINE = "DADDEB"
GREY = "6B7280"
INK = "1F2937"
WHITE = "FFFFFF"
GREEN = "15803D"
RED = "DC2626"

INR = "₹#,##0.00"
SIGNED_INR = "[Color10]+₹#,##0.00;[Red]-₹#,##0.00;₹0.00"  # gains green, losses red
QTY = "#,##0"
QTY_DEC = "#,##0.00"
SIGNED_QTY = "[Color10]+#,##0;[Red]-#,##0;0"
PCT = '0.00"%"'
DATE = "DD-MM-YYYY"

# status -> (fill, text)
STATUS_COLOURS = {
	"good": ("DCFCE7", "166534"),
	"wait": ("FEF3C7", "92400E"),
	"bad": ("FEE2E2", "991B1B"),
	"info": ("E0E7FF", "3730A3"),
	"none": ("F3F4F6", "374151"),
}
STATUS_KIND = {
	**dict.fromkeys(("Received", "Completed", "Return Completed", "Submitted", "Reconciled", "Paid", "Sale", "Debit Note", "No variance"), "good"),
	**dict.fromkeys(("Dispatched", "Partially Received", "Return Initiated", "Return Dispatched", "Draft", "Not created", "Pending",
		"Notes awaiting submission", "Unpaid", "Partly Paid", "Return", "To Bill", "To Deliver and Bill", "To Deliver", "Credit Note"), "wait"),
	**dict.fromkeys(("Cancelled", "Variance not invoiced", "Overdue"), "bad"),
	**dict.fromkeys(("Active", "Under Reconciliation", "Primary Sale", "Primary Return", "Closed", "Dispatch", "Return Delivery",
		"Secondary Sale", "Secondary Return", "Putaway", "Primary Invoice", "Sale Provision", "Return Credit Note"), "info"),
}

THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_TOP = Border(top=Side(style="medium", color=INDIGO), bottom=Side(style="thin", color=LINE))


def fill(colour):
	return PatternFill("solid", start_color=colour, end_color=colour)


@frappe.whitelist()
def download(name):
	"""Send the report workbook for one Primary Sales record."""
	from or_model.marketplace_or.ui_api import tracker_report

	data = tracker_report(name)  # checks read permission
	frappe.response["filename"] = f"{name} report {getdate().strftime('%Y-%m-%d')}.xlsx"
	frappe.response["filecontent"] = build(data)
	frappe.response["type"] = "binary"


def build(data):
	wb = Workbook()
	summary_sheet(wb.active, data)
	secondary_sheet(wb.create_sheet(), data)
	stock_sheet(wb.create_sheet(), data)
	accounting_sheet(wb.create_sheet(), data)
	notes_sheet(wb.create_sheet(), data)
	out = BytesIO()
	wb.save(out)
	return out.getvalue()


# ---------------------------------------------------------------- cell helpers


def _date(value):
	if not value:
		return None
	try:
		return getdate(value)
	except Exception:
		return value


def _qty_format(value, signed=False):
	if signed:
		return SIGNED_QTY if flt(value).is_integer() else "[Color10]+#,##0.00;[Red]-#,##0.00;0"
	return QTY if flt(value).is_integer() else QTY_DEC


def _link(doctype, name):
	return {"text": name, "link": get_url_to_form(doctype, name)} if name else None


def put(ws, row, col, value, kind="text", bold=False, colour=None, bg=None, size=11, align=None, border=True, italic=False):
	"""Write one cell. kind: text, date, qty, sqty, inr, sinr, pct, status, debit, credit."""
	cell = ws.cell(row=row, column=col)
	link = None
	if isinstance(value, dict):
		link, value = value.get("link"), value.get("text")
	numeric = kind in ("qty", "sqty", "inr", "sinr", "pct", "debit", "credit")
	if numeric and value in (None, ""):
		kind = "blank"  # nothing to show, not a zero
		value = None
	if kind == "date":
		value = _date(value)
		cell.number_format = DATE
	elif kind in ("qty", "sqty"):
		value = flt(value)
		cell.number_format = _qty_format(value, kind == "sqty")
	elif kind in ("inr", "sinr", "debit", "credit"):
		value = flt(value, 2)
		if kind in ("debit", "credit") and not value:
			value = None
		cell.number_format = SIGNED_INR if kind == "sinr" else INR
		if kind == "debit":
			colour = colour or GREEN
		elif kind == "credit":
			colour = colour or RED
	elif kind == "pct":
		value = flt(value)
		cell.number_format = PCT
	cell.value = value if value not in ("",) else None

	if kind == "status" and value:
		bg_colour, text_colour = STATUS_COLOURS[STATUS_KIND.get(str(value), "none")]
		bg, colour, bold = bg or bg_colour, text_colour, True
		align = align or "center"
	if link:
		cell.hyperlink = link
		colour, underline = colour or INDIGO, "single"
	else:
		underline = None
	cell.font = Font(name="Calibri", size=size, bold=bold, italic=italic, color=colour or INK, underline=underline)
	if bg:
		cell.fill = fill(bg)
	cell.alignment = Alignment(horizontal=align or ("right" if numeric else "left"), vertical="center", wrap_text=False, indent=0 if numeric else 1)
	if border:
		cell.border = BOX
	return cell


def merge(ws, row, col, last_col, value, bg=None, **kw):
	"""A value across several columns (the fill is painted on every merged cell)."""
	cell = put(ws, row, col, value, bg=bg, **kw)
	painted = cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type == "solid" else None
	if last_col > col:
		for c in range(col + 1, last_col + 1):
			other = ws.cell(row=row, column=c)
			if painted:  # the given colour, or a status pill's
				other.fill = fill(painted[-6:])
			if kw.get("border", True):
				other.border = BOX
		ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=last_col)
	return cell


def banner(ws, title, subtitle, width, tab):
	"""Rows 1-2: the sheet's title banner and a subtitle line."""
	ws.sheet_view.showGridLines = False
	ws.sheet_properties.tabColor = tab
	ws.row_dimensions[1].height = 34
	ws.row_dimensions[2].height = 20
	merge(ws, 1, 1, width, title, bg=INDIGO_DARK, colour=WHITE, bold=True, size=16, border=False)
	merge(ws, 2, 1, width, subtitle, bg=LAVENDER, colour=GREY, italic=True, size=10, border=False)
	ws.page_setup.orientation = "landscape"
	ws.page_setup.fitToWidth = 1
	ws.page_setup.fitToHeight = 0
	ws.sheet_properties.pageSetUpPr.fitToPage = True
	ws.print_title_rows = "1:2"


def section(ws, row, col, last_col, title):
	ws.row_dimensions[row].height = 22
	merge(ws, row, col, last_col, title.upper(), bg=SECTION, colour=INDIGO_DARK, bold=True, size=11, border=False)
	return row + 1


def table(ws, row, columns, rows, empty, totals=None, band_key=None, autofilter=False, freeze=False):
	"""A styled table at `row`: header, banded rows, optional totals row. columns: (title, kind). Returns the next row."""
	start = row
	ws.row_dimensions[row].height = 24
	for c, (title, kind) in enumerate(columns, start=1):
		put(ws, row, c, title, bold=True, colour=WHITE, bg=INDIGO, align="right" if kind in ("qty", "sqty", "inr", "sinr", "pct", "debit", "credit") else "left")
	if freeze:
		ws.freeze_panes = ws.cell(row=row + 1, column=1)
	row += 1
	if not rows:
		merge(ws, row, 1, len(columns), empty, colour=GREY, italic=True, align="center")
		return row + 2
	band, last_key = False, object()
	for i, values in enumerate(rows):
		if band_key:
			key = band_key(values)
			if key != last_key:
				band, last_key = not band, key
		else:
			band = i % 2 == 1
		for c, ((_title, kind), value) in enumerate(zip(columns, values), start=1):
			put(ws, row, c, value, kind=kind, bg=BAND if band and kind != "status" else None)
		row += 1
	if autofilter:
		ws.auto_filter.ref = f"A{start}:{get_column_letter(len(columns))}{row - 1}"
	if totals:
		for c, (_title, kind) in enumerate(columns, start=1):
			value = totals.get(c - 1, "")
			numeric = isinstance(value, (int, float)) and kind in ("qty", "sqty", "inr", "sinr", "pct", "debit", "credit")
			cell = put(ws, row, c, value, kind=kind if numeric else "text", bold=True, bg=LAVENDER)
			cell.border = TOTAL_TOP
		row += 1
	return row + 1


def widths(ws, values):
	for i, w in enumerate(values, start=1):
		ws.column_dimensions[get_column_letter(i)].width = w


def _subtitle(data):
	doc = data["doc"]
	return (
		f"{doc.get('marketplace')}  ·  PO {doc.get('marketplace_po_no')}  ·  Sales Order {doc.get('sales_order')}  ·  "
		f"Generated {str(data['generated_on'])[:16]} by {data['generated_by']}"
	)


# ---------------------------------------------------------------- sheets


def summary_sheet(ws, data):
	doc, so = data["doc"], data.get("sales_order") or {}
	ws.title = "Summary"
	widths(ws, [24, 16, 16, 16, 16, 16, 16, 16, 16, 16, 16, 16, 16, 16])
	banner(ws, f"Primary Sales Report  ·  {doc['name']}", _subtitle(data), 14, INDIGO)

	# figure tiles: label over a big value, four per row
	row = section(ws, 4, 1, 8, "Key figures")
	tiles = [
		[("Dispatched", doc.get("total_dispatched_qty"), "qty", "E0E7FF"), ("Received by marketplace", doc.get("marketplace_received_qty"), "qty", "DBEAFE"),
			("Stock at marketplace", doc.get("virtual_stock"), "qty", "DCFCE7"), ("Sold (secondary)", doc.get("total_sold_qty"), "qty", "FCE7F3")],
		[("Primary sales value", doc.get("primary_sales_value"), "inr", "E0E7FF"), ("Secondary sales value", doc.get("secondary_sales_value"), "inr", "DBEAFE"),
			("Marketplace commission", doc.get("marketplace_commission"), "inr", "FEF3C7"), ("Net price variance", doc.get("net_price_variance"), "sinr", "F3F4F6")],
	]
	for line in tiles:
		ws.row_dimensions[row].height = 20
		ws.row_dimensions[row + 1].height = 30
		for i, (label, value, kind, colour) in enumerate(line):
			col = 1 + i * 2
			merge(ws, row, col, col + 1, label, bg=colour, colour=GREY, size=10, bold=True, border=False)
			cell = merge(ws, row + 1, col, col + 1, value, kind=kind, bg=colour, size=18, bold=True, border=False, align="left")
			cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
		row += 3

	# record details | quantities, side by side
	top = row
	row = section(ws, top, 1, 4, "Primary sales record")
	details = [
		("Record", doc["name"], "text"), ("Marketplace", doc.get("marketplace"), "text"), ("Customer", so.get("customer_name") or doc.get("customer"), "text"),
		("Marketplace PO No", doc.get("marketplace_po_no"), "text"), ("PO date", doc.get("po_date"), "date"),
		("Sales Order", _link("Sales Order", doc.get("sales_order")), "text"), ("Order date", so.get("transaction_date"), "date"),
		("Order status", doc.get("order_status") or so.get("status"), "status"), ("OR stage", doc.get("status"), "status"),
		("Reconciliation", doc.get("reconciliation_status"), "status"), ("Approved by", doc.get("approved_by"), "text"),
	]
	for label, value, kind in details:
		put(ws, row, 1, label, bold=True, colour=GREY, bg=BAND)
		merge(ws, row, 2, 4, value, kind=kind, align="left" if kind != "status" else "center")
		row += 1
	left_end = row

	row = section(ws, top, 6, 8, "Quantities")
	quantities = [
		("Dispatched", doc.get("total_dispatched_qty")), ("Received by marketplace (POD)", doc.get("marketplace_received_qty")),
		("In transit", doc.get("stock_in_transit")), ("Sold (secondary)", doc.get("total_sold_qty")),
		("Secondary returns", doc.get("secondary_return_qty")), ("Stock at marketplace", doc.get("virtual_stock")),
		("Primary returns initiated", doc.get("return_initiated_qty")), ("Return in transit", doc.get("return_in_transit")),
		("Primary returns put away", doc.get("return_received_qty")), ("Dispatched, not invoiced", data.get("to_invoice")),
	]
	for label, value in quantities:
		merge(ws, row, 6, 7, label, bold=True, colour=GREY, bg=BAND)
		put(ws, row, 8, value, kind="qty", bold=True)
		row += 1
	row = max(row, left_end) + 1

	row = section(ws, row, 1, 4, "Values")
	values = [
		("Primary sales value (sold units)", doc.get("primary_sales_value"), "inr"), ("Secondary sales value", doc.get("secondary_sales_value"), "inr"),
		("Marketplace commission", doc.get("marketplace_commission"), "inr"), ("Net price variance", doc.get("net_price_variance"), "sinr"),
		("Debit notes (submitted)", doc.get("debit_note_amount"), "sinr"), ("Credit notes (submitted)", -flt(doc.get("credit_note_amount")), "sinr"),
		("OR Sale Provision (not yet credited)", -flt(doc.get("provision_amount")), "sinr"),
	]
	for label, value, kind in values:
		merge(ws, row, 1, 3, label, bold=True, colour=GREY, bg=BAND)
		put(ws, row, 4, value, kind=kind, bold=True)
		row += 1
	row += 1

	items = data.get("item_summary") or []
	row = section(ws, row, 1, 14, "Item summary")
	row = table(
		ws, row,
		[("Item", "text"), ("SKU", "text"), ("Dispatched", "qty"), ("Received", "qty"), ("In transit", "qty"), ("Sold", "qty"), ("Sec. returned", "qty"),
			("Returns dispatched", "qty"), ("At marketplace", "qty"), ("Last sale", "date"), ("Primary value", "inr"), ("Secondary value", "inr"),
			("Commission", "inr"), ("Net variance", "sinr")],
		[[i.get("item_code"), i.get("marketplace_sku"), i.get("dispatched"), i.get("received"), i.get("in_transit"), i.get("sold"), i.get("secondary_returned"),
			i.get("return_dispatched"), i.get("virtual_stock"), i.get("last_sale"), i.get("primary_value"), i.get("secondary_value"),
			i.get("commission_value"), i.get("net_variance")] for i in items],
		"No items yet",
		totals={0: "Total", **{k: sum(flt(i.get(f)) for i in items) for k, f in ((2, "dispatched"), (3, "received"), (4, "in_transit"), (5, "sold"),
			(6, "secondary_returned"), (7, "return_dispatched"), (8, "virtual_stock"), (10, "primary_value"), (11, "secondary_value"),
			(12, "commission_value"), (13, "net_variance"))}} if items else None,
	)

	moves = data["doc"].get("stock_movements") or []
	row = section(ws, row, 1, 14, "Primary sales & returns")
	row = table(
		ws, row,
		[("#", "qty"), ("Type", "status"), ("Item", "text"), ("Qty", "qty"), ("Received / put away", "qty"), ("Pending", "qty"), ("Status", "status"),
			("Return no", "text"), ("Delivery Note", "text"), ("Sales Invoice", "text"), ("Dispatch date", "date"), ("AWB", "text")],
		[[m.get("idx"), m.get("movement_type"), m.get("item_code"), m.get("qty"), m.get("received_qty"), m.get("pending_qty"), m.get("status"),
			m.get("reference_name") if m.get("movement_type") == "Primary Return" else "", _link("Delivery Note", m.get("delivery_note")),
			_link("Sales Invoice", m.get("sales_invoice")), m.get("dispatch_date"), m.get("awb_no")] for m in moves],
		"Nothing dispatched yet",
	)

	documents = (
		[[_link("Delivery Note", s.name), "Delivery Note", "Return Delivery" if s.is_return else "Dispatch", s.posting_date, s.qty, None,
			f"{s.warehouse or ''}{' · against ' + s.return_against if s.return_against else ''}{' · AWB ' + s.lr_no if s.lr_no else ''}"]
			for s in data.get("delivery_notes") or []]
		+ [[_link("Sales Invoice", s.name), "Sales Invoice", s.or_document_type, s.posting_date, None, s.grand_total,
			f"{s.customer or ''}{' · against ' + s.return_against if s.return_against else ''} · {s.status if s.docstatus else 'Draft'}"]
			for s in data.get("invoices") or []]
		+ [[_link("Journal Entry", j.name), "Journal Entry", j.or_purpose, j.posting_date, None, j.total_debit, j.or_against or ""]
			for j in data.get("journals") or []]
	)
	row = section(ws, row, 1, 14, "Documents created")
	table(
		ws, row,
		[("Document", "text"), ("Type", "text"), ("Step", "status"), ("Date", "date"), ("Qty", "qty"), ("Amount", "sinr"), ("Details", "text")],
		documents,
		"No documents yet",
	)
	ws.freeze_panes = "A3"


def secondary_sheet(ws, data):
	ws.title = "Secondary Sales"
	cols = [
		("Order / Return ID", "text"), ("Date", "date"), ("Type", "status"), ("Return against", "text"), ("SKU", "text"), ("Item", "text"),
		("Qty", "sqty"), ("Returned qty", "qty"), ("Primary rate", "inr"), ("Secondary rate", "inr"), ("Commission %", "pct"),
		("Rate w/o commission", "inr"), ("Secondary value", "sinr"), ("Variance / unit", "sinr"), ("Total variance", "sinr"),
		("Note status", "status"), ("Note", "text"), ("Primary invoice", "text"), ("Provision entry", "text"), ("Import batch", "text"),
	]
	widths(ws, [24, 12, 11, 18, 16, 16, 9, 12, 13, 14, 13, 17, 16, 15, 15, 14, 22, 22, 20, 20])
	banner(ws, f"Secondary Sales  ·  {data['doc']['name']}", _subtitle(data), len(cols), "16A34A")
	lines = data.get("lines") or []
	rows = [
		[l.marketplace_order_id, l.sale_date, "Return" if l.is_return else "Sale", l.return_against, l.marketplace_sku, l.item_code, l.qty, l.returned_qty,
			l.primary_rate, l.selling_rate, l.commission_percent, l.net_secondary_rate, flt(l.qty) * flt(l.selling_rate), l.variance_per_unit,
			l.total_variance, l.note_status, _link("Sales Invoice", l.sales_invoice), _link("Sales Invoice", l.primary_invoice),
			_link("Journal Entry", l.journal_entry), l.import_batch]
		for l in lines
	]
	table(
		ws, 4, cols, rows, "No secondary sales yet",
		totals={0: "Total", 6: sum(flt(l.qty) for l in lines), 12: sum(flt(l.qty) * flt(l.selling_rate) for l in lines),
			14: sum(flt(l.total_variance) for l in lines)} if lines else None,
		autofilter=bool(lines), freeze=True,
	)


def stock_sheet(ws, data):
	ws.title = "Stock Ledger"
	ledger = data.get("ledger") or {}
	stock = ledger.get("stock") or []
	cols = [("Date", "date"), ("Voucher type", "text"), ("Voucher", "text"), ("Item", "text"), ("Warehouse", "text"), ("Qty in / out", "sqty"),
		("Valuation rate", "inr"), ("Stock value change", "sinr"), ("Balance qty", "qty")]
	widths(ws, [13, 15, 22, 18, 28, 14, 15, 19, 13])
	banner(ws, f"Stock Ledger  ·  {data['doc']['name']}", _subtitle(data), len(cols), "0891B2")
	empty = "No access to the Stock Ledger" if ledger.get("can_see_stock") is False else "No stock movement yet"
	row = table(
		ws, 4, cols,
		[[s.posting_date, s.voucher_type, _link(s.voucher_type, s.voucher_no), s.item_code, s.warehouse, s.actual_qty, s.valuation_rate,
			s.stock_value_difference, s.qty_after_transaction] for s in stock],
		empty,
		totals={0: "Total", 7: sum(flt(s.stock_value_difference) for s in stock)} if stock else None,
		band_key=lambda values: values[2]["text"] if values[2] else None,  # one band per voucher
		autofilter=bool(stock), freeze=True,
	)
	if stock:
		net = defaultdict(lambda: [0.0, 0.0])
		for s in stock:
			net[(s.warehouse, s.item_code)][0] += flt(s.actual_qty)
			net[(s.warehouse, s.item_code)][1] += flt(s.stock_value_difference)
		row = section(ws, row, 1, 4, "Net position by warehouse (this record)")
		table(
			ws, row, [("Warehouse", "text"), ("Item", "text"), ("Net qty", "sqty"), ("Net value", "sinr")],
			[[w, item, q, v] for (w, item), (q, v) in sorted(net.items())], "",
		)


def accounting_sheet(ws, data):
	ws.title = "Accounting Ledger"
	ledger = data.get("ledger") or {}
	gl = ledger.get("gl") or []
	cols = [("Date", "date"), ("Voucher type", "text"), ("Voucher", "text"), ("Account", "text"), ("Party type", "text"), ("Party", "text"),
		("Debit", "debit"), ("Credit", "credit"), ("Cost center", "text")]
	widths(ws, [13, 16, 22, 34, 12, 24, 15, 15, 16])
	banner(ws, f"Accounting Ledger  ·  {data['doc']['name']}", _subtitle(data), len(cols), "7C3AED")
	empty = "No access to the General Ledger" if ledger.get("can_see_gl") is False else "No accounting entry yet"
	row = table(
		ws, 4, cols,
		[[g.posting_date, g.voucher_type, _link(g.voucher_type, g.voucher_no), g.account, g.party_type, g.party, g.debit, g.credit, g.cost_center]
			for g in gl],
		empty,
		totals={0: "Total", 6: sum(flt(g.debit) for g in gl), 7: sum(flt(g.credit) for g in gl)} if gl else None,
		band_key=lambda values: values[2]["text"] if values[2] else None,
		autofilter=bool(gl), freeze=True,
	)
	if gl:
		net = defaultdict(float)
		for g in gl:
			net[(g.account, g.party or "")] += flt(g.debit) - flt(g.credit)
		rows = [[acc, party, v if v > 0 else None, -v if v < 0 else None, v] for (acc, party), v in net.items() if abs(v) >= 0.005]
		row = section(ws, row, 1, 5, "Net effect by account (debit + / credit −)")
		table(
			ws, row, [("Account", "text"), ("Party", "text"), ("Debit", "debit"), ("Credit", "credit"), ("Net", "sinr")],
			rows, "Every account nets to zero",
		)


def notes_sheet(ws, data):
	ws.title = "Rate Difference Notes"
	notes = data.get("notes") or []
	cols = [("Note type", "status"), ("Note", "text"), ("Month · invoice", "text"), ("Lines", "qty"), ("Amount", "sinr"), ("Status", "status"),
		("Submitted on", "date"), ("Posting date", "date"), ("Customer", "text"), ("Against invoice", "text"), ("Grand total", "sinr")]
	widths(ws, [14, 24, 32, 9, 15, 13, 14, 14, 22, 24, 15])
	banner(ws, f"Rate Difference Notes  ·  {data['doc']['name']}", _subtitle(data), len(cols), "D97706")

	def signed(n):
		# a credit note lowers what the marketplace owes (red), a debit note raises it (green)
		return -flt(n.get("amount")) if n.get("note_type") == "Credit Note" else flt(n.get("amount"))

	rows = []
	for n in notes:
		inv = n.get("invoice") or {}
		rows.append([n.get("note_type"), _link("Sales Invoice", n.get("sales_invoice")), n.get("group_key"), n.get("line_count"), signed(n),
			n.get("status"), n.get("submitted_on"), inv.get("posting_date"), inv.get("customer"), _link("Sales Invoice", inv.get("return_against")),
			inv.get("grand_total")])
	row = table(
		ws, 4, cols, rows, "No rate difference notes yet",
		totals={0: "Net", 3: sum(flt(n.get("line_count")) for n in notes), 4: sum(signed(n) for n in notes)} if notes else None,
		autofilter=bool(notes), freeze=True,
	)
	if notes:
		row = section(ws, row, 1, 5, "By note type")
		by_type = defaultdict(lambda: [0, 0.0, 0.0])
		for n in notes:
			t = by_type[n.get("note_type")]
			t[0] += 1
			t[1] += signed(n)
			t[2] += signed(n) if n.get("status") == "Submitted" else 0
		table(
			ws, row, [("Note type", "status"), ("Notes", "qty"), ("Amount", "sinr"), ("Submitted", "sinr")],
			[[k, c, a, s] for k, (c, a, s) in sorted(by_type.items())], "",
		)
